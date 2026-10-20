import os, matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'); FIG = os.path.join(ROOT, 'figures')
# ---------- datasheet values (TI SN74AUP1G04, SCES571K, Rev. K, June 2014) ----------
TPD_TYP = {'1.2 V\u00b10.1': [5.9, 6.9, 7.8, 10.4], '1.5 V\u00b10.1': [4.2, 5.0, 5.6, 7.6], '1.8 V\u00b10.15': [3.4, 4.0, 4.6, 6.3], '2.5 V\u00b10.2': [2.5, 3.0, 3.5, 4.8], '3.3 V\u00b10.3': [2.1, 2.5, 2.9, 4.0]}
CLS = [5, 10, 15, 30]
fig, ax = plt.subplots(figsize=(6.4, 3.8))
for (k, v), c in zip(TPD_TYP.items(), ['#8e44ad', '#2980b9', '#16a085', '#e67e22', '#c0392b']): ax.plot(CLS, v, 'o-', color=c, label='V$_{CC}$ = ' + k)
ax.set_xlabel('C$_L$ (pF)'); ax.set_ylabel('t$_{pd}$ A\u2192Y, typical, T$_A$ = 25 \u00b0C (ns)'); ax.grid(alpha=.3); ax.legend(fontsize=8)
ax.set_title('SN74AUP1G04 typical propagation delay vs load (values read from datasheet Rev. K tables)', fontsize=8.5)
fig.tight_layout(); fig.savefig(os.path.join(FIG, 'fig_aup_tpd.png'), dpi=300); plt.close(fig)

# ---------- DOCX ----------
d = Document(); s = d.sections[0]; s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)
def H(t, l=1): h = d.add_heading(t, l); return h
def Pp(t, **k):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    for i, x in enumerate([t.replace('**','')]):
        if x: r = p.add_run(x); r.bold = k.get('bold', False); r.italic = k.get('italic', False)
    return p
def B(t):
    p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    for i, x in enumerate([t.replace('**','')]):
        if x: r = p.add_run(x); r.bold = False
def TB(rows, widths=None, size=8.5):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ''; r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size); r.bold = (i == 0)
            if widths: c.width = Inches(widths[j])
    d.add_paragraph().paragraph_format.space_after = Pt(2)
t = d.add_heading('Case Study: Texas Instruments SN74AUP1G04, Low-Power Single Inverter Gate', 0)
Pp('ECC-305 Microelectronics, Theory Case Study: commercially available IC / VLSI system (10 marks)', italic=True)
Pp('Student name / roll no. / batch / date: ____________________ (to be filled in)', italic=True)
Pp('**Convention used in this report.** "**[DS]**" = value or statement taken from the manufacturer\'s datasheet (SN74AUP1G04, document SCES571K, June 2004, revised June 2014, "Rev. K", retrieved from ti.com on 20 Oct 2026; TI\'s product page lists Rev. K as the current datasheet and the device status as ACTIVE). "**[Expl]**" = my own explanation or calculation, not a manufacturer specification.')
H('1. Device function and internal logic')
Pp('[DS] The SN74AUP1G04 is a single inverter gate performing the Boolean function Y = A\u0305 (function table: A = H \u2192 Y = L; A = L \u2192 Y = H). It belongs to TI\'s AUP ("Advanced Ultra-low-Power") family; per the datasheet the family has quiescent power consumption below 1 \u00b5A and the device is offered in several packages from SOT-23 (DBV, 2.90 \u00d7 1.60 mm nominal) down to the 0.64 mm\u00b2 X2SON (DPW, 0.8 \u00d7 0.8 mm) with 0.5 mm pitch. Pin functions (DBV/DCK/DRL): 1 N.C., 2 A (input), 3 GND, 4 Y (output), 5 V_CC.')
Pp('[DS] Features that influence the internal design: balanced outputs; input hysteresis (V_hys = 250 mV typ at 3.3 V, listed under Features) that tolerates slow input transitions; over-voltage-tolerant 3.6 V inputs/outputs; I_off circuitry for partial power-down and live insertion; reduced drive strength that "produces slower edges and prevents overshoot and undershoot".')
Pp('[Expl] The datasheet gives only a functional block diagram (an inverter symbol); the internal transistor netlist is not published. At the logic level the function is that of a static CMOS inverter: a PMOS pull-up and an NMOS pull-down driven by the same input (figure below). The real device also contains input hysteresis, Ioff/back-drive protection and (probably) buffer stages, so the generic circuit is illustrative only. This is the same structure that the BBM paper (Case Study 2, ECC-353) modifies to remove short-circuit current.')
d.add_picture(os.path.join(FIG, 'fig_cmos_inverter_generic.png'), width=Inches(2.6)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
H('2. Supply-voltage range')
TB([['Parameter [DS]', 'Value', 'Conditions / note'],
    ['Recommended V_CC', '0.8 V to 3.6 V', 'Recommended operating conditions (optimised for 3.3 V per Features)'],
    ['Absolute maximum V_CC', '\u22120.5 V to 4.6 V', 'Stress rating only'],
    ['Input voltage V_I', '0 to 3.6 V', 'Inputs 3.6 V tolerant at any valid V_CC; abs. max \u22120.5 to 4.6 V'],
    ['V_IH (min)', '0.65\u00d7V_CC (1.1\u20131.95 V); 1.6 V (2.3\u20132.7 V); 2 V (3\u20133.6 V)', 'High-level input voltage'],
    ['V_IL (max)', '0.35\u00d7V_CC (1.1\u20131.95 V); 0.7 V (2.3\u20132.7 V); 0.9 V (3\u20133.6 V)', 'Low-level input voltage'],
    ['Output drive I_OH / I_OL (max)', '\u00b120 \u00b5A (0.8 V); \u00b11.1 mA (1.1 V); \u00b11.7 mA (1.4 V); \u00b11.9 mA (1.65 V); \u00b13.1 mA (2.3 V); \u00b14 mA (3 V)', 'Recommended operating conditions'],
    ['Input transition rate \u0394t/\u0394v', '200 ns/V max', 'V_CC = 0.8 to 3.6 V'],
    ['Operating temperature', '\u221240 \u00b0C to 85 \u00b0C', 'Free-air']], widths=[1.8, 2.7, 2.2])
H('3. Power-related specifications')
TB([['Parameter [DS]', 'Value', 'Test conditions'],
    ['Static supply current I_CC', '0.5 \u00b5A max at 25 \u00b0C; 0.9 \u00b5A max over \u221240\u201385 \u00b0C (the 0.9 \u00b5A figure is also quoted in Features)', 'V_I = GND or V_CC, I_O = 0, V_CC = 0.8 to 3.6 V'],
    ['\u0394I_CC (one input held below rail)', '40 \u00b5A (25 \u00b0C) / 50 \u00b5A (\u221240\u201385 \u00b0C) max', 'V_I = V_CC \u2212 0.6 V, I_O = 0, V_CC = 3.3 V'],
    ['I_off (partial power-down)', '0.2 \u00b5A (25 \u00b0C) / 0.6 \u00b5A (\u221240\u201385 \u00b0C) max', 'V_CC = 0 V, V_I or V_O = 0 to 3.6 V'],
    ['Input leakage I_I', '0.1 \u00b5A (25 \u00b0C) / 0.5 \u00b5A (\u221240\u201385 \u00b0C) max', 'V_I = GND to 3.6 V, V_CC = 0 to 3.6 V'],
    ['Power dissipation capacitance C_pd', '3.9 pF typ (V_CC = 0.8 to 2.5 V); 4.1 pF typ (3.3 V \u00b1 0.3 V)', 'f = 10 MHz, T_A = 25 \u00b0C'],
    ['Input capacitance C_i / output C_o', '1.5 pF typ / 2.5 pF typ', 'V_I = V_CC or GND; C_o at V_CC = 3.6 V'],
    ['Output levels', 'V_OH \u2265 V_CC \u2212 0.1 V and V_OL \u2264 0.1 V at |I_O| = 20 \u00b5A (0.8\u20133.6 V); e.g. V_OH \u2265 1.32 V (25 \u00b0C) at 1.65 V, I_OH = \u22121.9 mA', 'See datasheet Section 7.5'],
    ['Thermal resistance R\u03b8JA', '298.6 \u00b0C/W (DBV) \u2026 554.9 \u00b0C/W (DRY)', 'Thermal information table'],
    ['ESD', 'HBM 2000 V, CDM 1000 V', 'Handling ratings; latch-up > 100 mA per JESD 78 Class II']], widths=[1.8, 2.8, 2.1])
Pp('Note on column assignment: the I_CC, \u0394I_CC, I_off and I_I rows list only two numbers in the text extraction of the PDF; they are read here as (T_A = 25 \u00b0C max, \u221240\u201385 \u00b0C max), which is consistent with the Features statement "I_CC = 0.9 \u00b5A max". The reader should confirm against the PDF table if exact column placement matters.', italic=True)
Pp('One discrepancy turned up while checking. TI\'s product page parametric table lists "Supply current (max) 10 \u00b5A", while the datasheet\'s electrical table gives I_CC = 0.9 \u00b5A max (and \u0394I_CC up to 50 \u00b5A). The parametric figure is probably defined under a different convention (e.g. a worst-case across the whole family); this report uses the datasheet value. The product page also gives the typical-delay descriptor "tpd 5\u201310 ns", which is a family-level figure, not the specification at a given V_CC and load.')
H('4. Propagation delay and loading conditions')
Pp('[DS] Test conditions: load circuit of datasheet Figure 3 with C_L = 5, 10, 15 or 30 pF (including probe and jig), 1 M\u03a9 resistor to the output, measured at V_M = V_CC/2; input pulses from a generator with PRR \u2264 10 MHz, Z_O = 50 \u03a9, t_r/t_f \u2264 3 ns; outputs measured one at a time, t_PLH and t_PHL are the same as t_pd.')
TB([['t_pd A\u2192Y [DS]', 'T_A = 25 \u00b0C: min / typ / max (ns)', '\u221240 \u00b0C to 85 \u00b0C: min / max (ns)'],
    ['C_L = 5 pF, V_CC = 1.8 V \u00b1 0.15 V', '2.2 / 3.4 / 5.9', '1.4 / 7.0'],
    ['C_L = 5 pF, V_CC = 3.3 V \u00b1 0.3 V', '1.4 / 2.1 / 3.2', '1.2 / 3.9'],
    ['C_L = 5 pF, V_CC = 1.2 V \u00b1 0.1 V', '3.3 / 5.9 / 10.8', '2.1 / 13.5'],
    ['C_L = 10 pF, V_CC = 1.8 V \u00b1 0.15 V', '2.6 / 4.0 / 6.9', '2.1 / 7.9'],
    ['C_L = 15 pF, V_CC = 1.8 V \u00b1 0.15 V', '3.2 / 4.6 / 7.4', '2.6 / 8.5'],
    ['C_L = 30 pF, V_CC = 1.8 V \u00b1 0.15 V', '4.6 / 6.3 / 9.2', '4.1 / 10.5'],
    ['C_L = 30 pF, V_CC = 3.3 V \u00b1 0.3 V', '3.2 / 4.0 / 5.4', '2.9 / 6.2']], widths=[2.6, 2.1, 2.0])
Pp('Other values from the same tables: at V_CC = 0.8 V the datasheet gives a single t_pd value of 15.6 ns (C_L = 5 pF), 17.7 ns (10 pF), 19.5 ns (15 pF) and 25.4 ns (30 pF). Features: t_pd = 3.9 ns max at 3.3 V (this is the C_L = 5 pF, \u221240\u201385 \u00b0C maximum).')
d.add_picture(os.path.join(FIG, 'fig_aup_tpd.png'), width=Inches(5.0)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
Pp('Delay rises with load (about +0.5 ns per 5 pF near 1.8 V for the first steps, then more slowly per pF) and rises steeply as V_CC falls toward 0.8 V, because the drive current falls roughly with (V_CC \u2212 V_T)\u00b2 [Expl].', italic=True)
H('5. Worked example: dynamic versus static power  [Expl, calculation using datasheet C_pd, C_L, I_CC]')
Pp('Using the standard definition P_dyn = (C_pd + C_L) \u00b7 V_CC\u00b2 \u00b7 f (a general formula, not a datasheet statement) with V_CC = 1.8 V, f = 10 MHz, C_pd = 3.9 pF (datasheet), C_L = 5 pF (the datasheet test load): P_dyn = (3.9 pF + 5 pF) \u00d7 (1.8 V)\u00b2 \u00d7 10 MHz = 8.9\u00d710\u207b\u00b9\u00b2 \u00d7 3.24 \u00d7 10\u2077 \u2248 **288 \u00b5W**. The worst-case static power from the datasheet is I_CC(max) \u00d7 V_CC = 0.9 \u00b5A \u00d7 1.8 V \u2248 **1.6 \u00b5W**, i.e. under 1 % of the dynamic power at 10 MHz. Dynamic power therefore dominates whenever the gate toggles; static power matters mainly for rarely toggling, battery-backed nodes. These are calculations at one chosen operating point, not specifications.')
H('6. Applications and engineering relevance')
B('[DS] Listed applications include ATCA solutions, active noise cancellation, barcode scanners, blood-pressure monitors, CPAP machines, cable solutions, embedded PCs, field transmitters, fingerprint biometrics, HVAC, NAS, server motherboards/PSU, software-defined radio, TVs, video-communications systems, wireless data-access cards/headsets/keyboards/mice and X-ray equipment.')
B('[DS] Section 10: the AUP family is positioned for battery-powered portable products; the 3.6 V-tolerant I/O allows down-translation, e.g. dropping 3.3 V peripheral drivers to the 0.8\u20131.2 V levels of newer microcontrollers.')
B('[Expl] Typical uses: signal inversion, glue logic, clock/enable conditioning, level-compatible interfacing, and (as a building block) simple oscillators/delay elements when combined with R and C. Relevance to ECC-305: a real-world instance of low-power CMOS design, small C_pd and I_CC, reduced drive and slow edges to cut dynamic power and ground bounce, hysteresis for slow inputs, and I_off for partial power-down. The BBM paper (PBL component) is about removing the short-circuit part of dynamic power in exactly this kind of inverter-based stage.')
H('7. Limitations and design trade-offs')
B('**Speed vs. power:** t_pd is a few ns even at 3.3 V (3.2 ns max at 25 °C, 3.9 ns over temperature, C_L = 5 pF) and 10\u201315 ns at 1.2 V\u20130.8 V; the datasheet does not quote a maximum clock frequency.')
B('**Drive strength:** only \u00b14 mA at 3 V, \u00b11.9 mA at 1.65 V and \u00b120 \u00b5A at 0.8 V; unsuitable for driving LEDs, long unterminated lines or heavy capacitive loads (delay rises with C_L).')
B('**Input conditions:** input rise/fall must not exceed 200 ns/V; unused inputs must be tied to V_CC or GND; an input held between rails (\u0394I_CC condition) can draw up to 50 \u00b5A at 3.3 V.')
B('**Absolute limits:** V_CC \u2264 4.6 V (abs. max), continuous output current \u00b120 mA, V_CC/GND current \u00b150 mA; ESD ratings are modest (HBM 2000 V, CDM 1000 V), so ESD handling is needed.')
B('**Thermal:** R\u03b8JA ~300\u2013555 \u00b0C/W depending on package; irrelevant at \u00b5W power but limits sustained heavy drive.')
B('**Family trade-offs (TI product page, Similar products):** SN74LVC1G04 offers 1.65\u20135.5 V operation and much higher drive (stated average 24 mA) at the cost of higher power; TI also lists the newer SN74UVP1G04 (0.65\u20133.6 V, preview) as a lower-voltage alternative. [DS/TI page]')
H('8. Questions the audience may ask (technically accurate answers)')
QA = [('What does Y = A\u0305 mean and is the output ever undefined?', 'It is a logical inversion. Provided V_I is within the recommended levels (V_IL max / V_IH min for the V_CC used) and V_CC is within 0.8\u20133.6 V, the output is a defined logic level; between V_IL and V_IH the output is not guaranteed, which is why slow or floating inputs must be avoided (hysteresis helps with slow edges).'),
 ('Why is dynamic power much larger than static power?', 'Static current is only leakage (I_CC \u2264 0.9 \u00b5A), whereas every transition charges and discharges C_pd and C_L: P = (C_pd + C_L)V_CC\u00b2f. At 10 MHz, 1.8 V, 5 pF this is about 288 \u00b5W versus under 2 \u00b5W static (calculation above).'),
 ('How does the load capacitance change the delay?', 'Datasheet at 1.8 V, 25 \u00b0C typ: 3.4 ns (5 pF), 4.0 ns (10 pF), 4.6 ns (15 pF), 6.3 ns (30 pF), delay increases roughly linearly with C_L because the weak output stage supplies limited current.'),
 ('Why does the speed drop at low V_CC?', 'Drive current falls strongly as V_CC approaches V_T; typical t_pd rises from 2.1 ns at 3.3 V to 5.9 ns at 1.2 V (C_L = 5 pF), and a single 15.6 ns value is quoted at 0.8 V.'),
 ('What is I_off for?', 'With V_CC = 0 V the I_off circuitry disables the outputs and limits leakage to \u2264 0.6 \u00b5A (V_I or V_O up to 3.6 V), preventing back-drive current, which allows live insertion and partial power-down.'),
 ('What is the purpose of input hysteresis?', 'The datasheet states V_hys = 250 mV typ at 3.3 V; it gives better noise immunity and allows slow input transitions without output chatter. It does not remove the 200 ns/V limit.'),
 ('Can a 3.3 V signal drive this gate when V_CC = 1.2 V?', 'Yes: inputs and outputs are specified 3.6 V tolerant, so the gate can down-translate 3.3 V signals to its V_CC level (datasheet applications section). Output levels are then limited to V_CC.'),
 ('How is this related to break-before-make?', 'Any static CMOS inverter has a brief moment of PMOS and NMOS both conducting (short-circuit current). AUP devices reduce it through slower, controlled edges and low drive; the BBM technique in the research paper removes it structurally with time-skewed gate signals but costs extra transistors and wiring. TI does not state that the AUP1G04 uses BBM.'),
 ('Why choose this gate instead of SN74LVC1G04?', 'For battery-powered products where power and tiny packages matter and loads are small; LVC offers higher drive and 5.5 V operation but consumes more power.'),
 ('Which numbers should I quote if asked for "the" propagation delay?', 'Always give V_CC, C_L, temperature and whether it is min/typ/max: e.g. 3.4 ns typ at 1.8 V, 5 pF, 25 \u00b0C; 3.9 ns max at 3.3 V over temperature.')]
for q, a in QA: Pp('**Q: ' + q + '**'); Pp('A: ' + a)
H('9. References')
for r in ['[1] Texas Instruments, "SN74AUP1G04 Low-Power Single Inverter Gate," datasheet SCES571K, June 2004, revised June 2014 (Rev. K), https://www.ti.com/lit/gpn/SN74AUP1G04 (retrieved 20 Oct 2026; package addendum dated Oct 2025).',
          '[2] Texas Instruments, SN74AUP1G04 product page, https://www.ti.com/product/SN74AUP1G04 (status ACTIVE; parametric data; similar products; retrieved 20 Oct 2026).',
          '[3] J. Puhan, D. Raič, T. Tuma, \u00c1. B\u0171rmen, "Break-before-Make CMOS Inverter for Power-Efficient Delay Implementation," The Scientific World Journal, vol. 2014, Art. 349131, doi:10.1155/2014/349131.']: Pp(r)
d.save(os.path.join(ROOT, 'Case_Study.docx')); print('docx ok')
